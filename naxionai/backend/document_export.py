"""
NAXIONAI — Markdown -> .docx / PDF export (Phase 6).

Ported from neuroworks/server/src/lib/docx-export.ts + the PDF half of
lib/browser.ts. Markdown is rendered to HTML (the `markdown` package) and
walked with BeautifulSoup — both already naxionai dependencies or, in
markdown's case, a small pure-Python addition — into python-docx primitives
for .docx, or handed to the existing Ghost Browser (Playwright) for a
print-to-PDF pass, reusing infrastructure this app already has instead of
adding a headless-browser dependency of its own.
"""

from io import BytesIO

from bs4 import BeautifulSoup, Tag
from docx import Document
from docx.shared import Pt
import markdown as md


def _add_inline(paragraph, node) -> None:
    """Recursively render inline HTML (bold/italic/code/links) as runs on
    an existing docx paragraph."""
    if isinstance(node, str):
        if node:
            paragraph.add_run(node)
        return
    if node.name == "strong" or node.name == "b":
        for child in node.children:
            run_text = child if isinstance(child, str) else child.get_text()
            r = paragraph.add_run(run_text)
            r.bold = True
    elif node.name == "em" or node.name == "i":
        for child in node.children:
            run_text = child if isinstance(child, str) else child.get_text()
            r = paragraph.add_run(run_text)
            r.italic = True
    elif node.name == "code":
        r = paragraph.add_run(node.get_text())
        r.font.name = "Consolas"
    elif node.name == "a":
        paragraph.add_run(node.get_text())
    else:
        for child in node.children:
            _add_inline(paragraph, child)


_HEADING_STYLE = {f"h{i}": f"Heading {i}" for i in range(1, 7)}


def markdown_to_docx_bytes(markdown_text: str, title: str = "Document") -> bytes:
    html = md.markdown(markdown_text or "", extensions=["tables", "fenced_code"])
    soup = BeautifulSoup(html, "html.parser")

    doc = Document()
    doc.core_properties.title = title

    for el in soup.find_all(recursive=False):
        _render_block(doc, el)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _render_block(doc: Document, el: Tag) -> None:
    if el.name in _HEADING_STYLE:
        p = doc.add_paragraph(style=_HEADING_STYLE[el.name])
        _add_inline(p, el)
    elif el.name == "p":
        p = doc.add_paragraph()
        for child in el.children:
            _add_inline(p, child)
    elif el.name == "blockquote":
        for inner in el.find_all(["p"], recursive=False) or [el]:
            p = doc.add_paragraph(style="Intense Quote")
            _add_inline(p, inner)
    elif el.name in ("ul", "ol"):
        style = "List Bullet" if el.name == "ul" else "List Number"
        for li in el.find_all("li", recursive=False):
            p = doc.add_paragraph(style=style)
            for child in li.children:
                if getattr(child, "name", None) in ("ul", "ol"):
                    continue  # nested lists: flatten by skipping (rare in agent-generated reports)
                _add_inline(p, child)
    elif el.name == "pre":
        code = el.get_text()
        for line in code.split("\n"):
            p = doc.add_paragraph()
            r = p.add_run(line)
            r.font.name = "Consolas"
            r.font.size = Pt(9)
    elif el.name == "table":
        rows = el.find_all("tr")
        if not rows:
            return
        n_cols = len(rows[0].find_all(["th", "td"]))
        table = doc.add_table(rows=0, cols=n_cols)
        table.style = "Light Grid Accent 1"
        for tr in rows:
            cells = tr.find_all(["th", "td"])
            row = table.add_row()
            for i, cell in enumerate(cells):
                row.cells[i].text = cell.get_text().strip()
    elif el.name == "hr":
        doc.add_paragraph("_" * 40)
    elif el.name is not None:
        text = el.get_text().strip()
        if text:
            doc.add_paragraph(text)


async def markdown_to_pdf_bytes(markdown_text: str, title: str = "Document") -> bytes:
    """Renders markdown -> HTML -> PDF using the Ghost Browser's persistent
    Playwright context (already running for scraping) instead of adding a
    separate headless-browser dependency."""
    from agent.browser.ghost import ghost

    if not ghost.is_available:
        raise RuntimeError("PDF export unavailable: Ghost Browser is offline (see startup logs)")

    html_body = md.markdown(markdown_text or "", extensions=["tables", "fenced_code"])
    full_html = f"""<!doctype html><html><head><meta charset="utf-8">
<title>{title}</title>
<style>
body {{ font-family: Calibri, Arial, sans-serif; font-size: 12pt; margin: 2.5cm; }}
pre, code {{ font-family: Consolas, monospace; background: #f5f5f5; }}
table {{ border-collapse: collapse; width: 100%; }}
th, td {{ border: 1px solid #999; padding: 4px 8px; }}
</style></head><body>{html_body}</body></html>"""

    page = await ghost.get_page("_document_export")
    try:
        await page.set_content(full_html, wait_until="load")
        return await page.pdf(format="A4", margin={"top": "1cm", "bottom": "1cm", "left": "1cm", "right": "1cm"})
    finally:
        await page.close()
